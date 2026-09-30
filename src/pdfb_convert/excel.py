# SPDX-License-Identifier: AGPL-3.0-only

"""PDF to Excel (spec §4.5): each page's tables go to separate worksheets; the text lines of all pages, in page order,
go to the last worksheet, "Metin" (Turkish for "Text"). The "Metin" sheet is written for documents with tables as well
(controller decision, 2c): borderless tables and text outside tables are not found by find_tables and would otherwise
be lost. The tables count covers tables only (the API adds its NO_TABLES note when tables == 0).

Turkish numbers and dates are converted to real cell types by tr_numbers; everything else is a text cell (xlsx_safe).
The file name, the document metadata and cell contents are never written to stdout, to the log or to the output's
properties.

Error contract (same as word.py): opening and limits are handled in open_pdf (11, 12, 13); a scanned PDF is 10. A
MuPDF error raised after opening, while pages are read or tables are extracted, means an unreadable file (12). A
failing page is not skipped: every non-MuPDF error propagates and the CLI turns it into 1 (CONVERT_FAILED); no
incomplete workbook is written.
"""

import logging
from collections.abc import Iterator, Sequence
from contextlib import contextmanager

import pymupdf
from openpyxl import Workbook
from openpyxl.packaging import extended
from openpyxl.worksheet.worksheet import Worksheet

from pdfb_convert.errors import ConvertError, ScannedPdf, Unreadable
from pdfb_convert.pdf import open_pdf
from pdfb_convert.scan import is_scanned, page_char_counts
from pdfb_convert.tr_numbers import column_has_comma_decimal, parse_cell
from pdfb_convert.xlsx_safe import put_cell, sheet_title

TEXT_SHEET = "Metin"
# Excel's worksheet row limit. If the "Metin" sheet would exceed it, the document is rejected before any writing
# (review M7).
MAX_ROWS = 1_048_576
# Internal log: numbers only (the count of truncated cells). The CLI disables logging and points stderr at /dev/null;
# this record is not part of the stdout contract ({"pages","tables"}) and exists for tests and future internal metrics
# (review M6).
_log = logging.getLogger(__name__)
# MuPDF errors raised after opening. RuntimeError/ValueError from pdf.MUPDF_ERRORS are deliberately absent: errors in
# the Python layer of the table extractor are conversion failures (CONVERT_FAILED), not unreadable files.
_LATE_MUPDF_ERRORS = (pymupdf.FileDataError, pymupdf.mupdf.FzErrorBase)
# openpyxl's DocumentProperties carry the creator "openpyxl" and empty fields; all are cleared (cf. word.py).
_EMPTY_PROPS = (
    "creator",
    "title",
    "subject",
    "description",
    "keywords",
    "category",
    "identifier",
    "lastModifiedBy",
    "contentStatus",
    "language",
    "version",
)


class TooManyRows(ConvertError):
    """The text lines exceed Excel's row limit (review M7). Exit code 1: CONVERT_FAILED.

    TOO_COMPLEX was not chosen: in spec §4.2 that code is for the upload's archive limits (§4.3, sniff stage) and it has
    no counterpart in the Python exit contract (10-13, everything else CONVERT_FAILED; python.ts PYTHON_EXIT). A new
    code would change the API mapping, the spec table and the UI text; a million lines in 100 pages only happens with a
    pathological PDF.
    """

    exit_code = 1


def _cell_text(v: object) -> str | None:
    if v is None:
        return None
    return " ".join(str(v).split())  # line breaks inside a cell become a single space


def _write_table(ws: Worksheet, rows: Sequence[Sequence[object]]) -> int:
    width = max((len(r) for r in rows), default=0)
    columns = [[_cell_text(r[c]) or "" for r in rows if c < len(r)] for c in range(width)]
    comma = [column_has_comma_decimal(col) for col in columns]
    truncated = 0
    for ri, row in enumerate(rows, start=1):
        for ci, raw in enumerate(row, start=1):
            text = _cell_text(raw)
            if not text:
                continue  # continuation of a merged cell, or empty: best effort
            truncated += put_cell(ws, ri, ci, text, parse_cell(text, comma[ci - 1]))
    return truncated


def _write_tables(doc: pymupdf.Document, wb: Workbook, used: set[str]) -> tuple[int, int]:
    tables = truncated = 0
    for pno in range(doc.page_count):
        page = doc.load_page(pno)
        tno = 0
        for table in page.find_tables().tables:
            rows = table.extract()
            if not any(any(c for c in r) for r in rows):
                continue
            tno += 1
            tables += 1
            # Worksheet names are user-facing Turkish: "Sayfa N - Tablo M" means "Page N - Table M".
            truncated += _write_table(wb.create_sheet(sheet_title(f"Sayfa {pno + 1} - Tablo {tno}", used)), rows)
    return tables, truncated


def _text_lines(doc: pymupdf.Document) -> list[str]:
    """Non-empty text lines of all pages, in page order. Raises TooManyRows up front (before any table or worksheet is
    written) if the row limit would be exceeded."""
    lines = [
        t for pno in range(doc.page_count) for t in map(_cell_text, doc.load_page(pno).get_text("text").splitlines())
    ]
    lines = [t for t in lines if t]
    if len(lines) > MAX_ROWS:
        raise TooManyRows()
    return lines


def _write_text(lines: list[str], wb: Workbook, used: set[str]) -> int:
    ws = wb.create_sheet(sheet_title(TEXT_SHEET, used))
    # No number conversion on the text sheet.
    return sum(put_cell(ws, r, 1, text, None) for r, text in enumerate(lines, start=1))


def _neutral_workbook() -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    for field in _EMPTY_PROPS:
        setattr(wb.properties, field, None)
    return wb


@contextmanager
def _neutral_app_properties() -> Iterator[None]:
    """Keeps the server software and its version out of app.xml (review M8): Excel's own value, no AppVersion.

    openpyxl's writer takes ExtendedProperties from this module while saving and writes a fixed Application. A
    subclass does not work (Serialisable builds an empty element list for subclasses), so a factory is installed for
    the duration of the save and the original class is restored afterwards (single-threaded, one-job process).
    """
    original = extended.ExtendedProperties

    def neutral() -> extended.ExtendedProperties:
        props = original()
        props.Application = "Microsoft Excel"
        props.AppVersion = None
        return props

    extended.ExtendedProperties = neutral
    try:
        yield
    finally:
        extended.ExtendedProperties = original


def convert_excel(src: str, dst: str, max_pages: int) -> dict[str, int]:
    doc = open_pdf(src, max_pages)
    try:
        try:
            counts = page_char_counts(doc)
        except MemoryError:
            raise
        except Exception as exc:  # noqa: BLE001 - MuPDF errors such as a page tree cycle are raised while reading pages
            raise Unreadable() from exc
        if is_scanned(counts):
            raise ScannedPdf()
        wb = _neutral_workbook()
        used: set[str] = set()
        try:
            lines = _text_lines(doc)
            tables, truncated = _write_tables(doc, wb, used)
            truncated += _write_text(lines, wb, used)  # always, and always last
        except _LATE_MUPDF_ERRORS as exc:
            raise Unreadable() from exc
        if truncated:
            _log.info("cells truncated at 32767 characters: %d", truncated)
        with _neutral_app_properties():
            wb.save(dst)
        return {"pages": doc.page_count, "tables": tables}
    finally:
        doc.close()
