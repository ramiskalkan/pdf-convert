# SPDX-License-Identifier: AGPL-3.0-only

import json
import os
import subprocess
import sys
import zipfile
from datetime import datetime
from pathlib import Path

import pymupdf
import pytest
from openpyxl import load_workbook

from conftest import BROKEN_TREES, image_only_pdf
from corpus.make_corpus import make_excel_invoice, make_excel_no_tables
from pdfb_convert import excel as excel_mod
from pdfb_convert import xlsx_safe
from pdfb_convert.errors import ConvertError, EncryptedPdf, ScannedPdf, TooManyPages, Unreadable
from pdfb_convert.excel import TooManyRows, convert_excel

SRC = str(Path(__file__).resolve().parents[1] / "src")


@pytest.fixture()
def invoice(tmp_path: Path, tr_font: str) -> Path:
    return make_excel_invoice(tmp_path / "secret-customer-name.pdf", tr_font)


def test_tables_on_separate_sheets_and_numbers_typed(invoice: Path, tmp_path: Path):
    out = tmp_path / "output.xlsx"
    assert convert_excel(str(invoice), str(out), 100) == {"pages": 2, "tables": 3}
    wb = load_workbook(out)
    # Controller decision (2c): the "Metin" sheet also exists in a document with tables, and it comes last.
    assert wb.sheetnames == ["Sayfa 1 - Tablo 1", "Sayfa 1 - Tablo 2", "Sayfa 2 - Tablo 1", "Metin"]
    ws = wb["Sayfa 1 - Tablo 1"]
    rows = {ws.cell(row=i, column=1).value: i for i in range(1, ws.max_row + 1)}
    r2 = rows["Danışmanlık hizmeti"]
    # The quantity column has no comma decimal: the ambiguous 1.250 stays text (column rule).
    assert ws.cell(row=r2, column=2).value == "1.250" and ws.cell(row=r2, column=2).data_type == "s"
    assert ws.cell(row=r2, column=3).value == pytest.approx(1234.56)
    assert ws.cell(row=r2, column=3).number_format == '"₺"#,##0.00'
    assert ws.cell(row=r2, column=4).value == pytest.approx(1543200.0)
    assert isinstance(ws.cell(row=r2, column=5).value, datetime)
    assert ws.cell(row=rows["KDV oranı"], column=3).value == pytest.approx(0.20)
    # Formula injection row: all text
    rf = rows["=TOPLA(D2:D4)"]
    for col in range(1, 6):
        c = ws.cell(row=rf, column=col)
        assert c.data_type == "s", (col, c.value)
    assert ws.cell(row=rf, column=3).value == "00123"
    assert ws.cell(row=rf, column=5).value == "31.02.2026"


def test_ambiguous_number_column_rule(invoice: Path, tmp_path: Path):
    out = tmp_path / "o.xlsx"
    convert_excel(str(invoice), str(out), 100)
    ws = load_workbook(out)["Sayfa 1 - Tablo 2"]
    vals = {ws.cell(row=i, column=1).value: ws.cell(row=i, column=2).value for i in range(1, ws.max_row + 1)}
    # The column contains 12.500,00 (comma decimal): 1.234 is read as a thousands separator.
    assert vals["Şubat"] == 1234
    assert vals["Ocak"] == pytest.approx(12500.0)


def test_no_cell_is_a_formula(invoice: Path, tmp_path: Path):
    out = tmp_path / "o.xlsx"
    convert_excel(str(invoice), str(out), 100)
    with zipfile.ZipFile(out) as z:
        for name in z.namelist():
            if name.startswith("xl/worksheets/"):
                assert b"<f" not in z.read(name), name


def test_text_sheet_when_no_tables(tmp_path: Path, tr_font: str):
    src = make_excel_no_tables(tmp_path / "t.pdf", tr_font)
    out = tmp_path / "o.xlsx"
    assert convert_excel(str(src), str(out), 100) == {"pages": 2, "tables": 0}
    wb = load_workbook(out)
    assert wb.sheetnames == ["Metin"]
    col = [c.value for c in wb["Metin"]["A"] if c.value is not None]
    assert col[0] == "Dilekçe sayfa 1"
    assert "Dilekçe sayfa 2" in col
    assert all(wb["Metin"].cell(row=i, column=1).data_type == "s" for i in range(1, len(col) + 1))
    assert "1.234,56" in col  # no number conversion on the text sheet


def _text_sheet(out: Path) -> list[str]:
    return [c.value for c in load_workbook(out)["Metin"]["A"] if c.value is not None]


def test_text_sheet_is_last_and_holds_all_text_in_document_with_tables(invoice: Path, tmp_path: Path):
    # Controller decision (2c): text outside ruled tables is not lost; "Metin" holds the text of every page in page
    # order. The tables count (and the API's NO_TABLES note) looks at tables only.
    out = tmp_path / "o.xlsx"
    assert convert_excel(str(invoice), str(out), 100) == {"pages": 2, "tables": 3}
    wb = load_workbook(out)
    assert wb.sheetnames[-1] == "Metin"
    col = _text_sheet(out)
    title1 = col.index("Fatura dökümü — Çağ Ltd. Şti.")
    title2 = col.index("Ek tablo: bordro özeti")
    assert title1 < col.index("Danışmanlık hizmeti") < title2 < col.index("Ayşe Yılmaz")
    ws = wb["Metin"]
    assert all(ws.cell(row=i, column=1).data_type == "s" for i in range(1, len(col) + 1))
    # No number/date conversion on the text sheet; a line that looks like a formula is text.
    assert "12.500,00" in col and "=TOPLA(D2:D4)" in col
    assert ws.cell(row=col.index("=TOPLA(D2:D4)") + 1, column=1).data_type == "s"


def test_borderless_table_lands_on_text_sheet(invoice: Path, tmp_path: Path, tr_font: str):
    # find_tables does not detect a borderless table; its cells land on the "Metin" sheet and are not lost.
    doc = pymupdf.open(invoice)
    page = doc.new_page()
    page.insert_font(fontname="tr", fontfile=tr_font)
    for i, (item, qty) in enumerate([("Ürün", "Adet"), ("Kurşun kalem", "12"), ("Silgi", "3")]):
        page.insert_text((60, 80 + 18 * i), item, fontname="tr", fontsize=11)
        page.insert_text((300, 80 + 18 * i), qty, fontname="tr", fontsize=11)
    src = tmp_path / "borderless.pdf"
    doc.save(src)
    doc.close()
    out = tmp_path / "o.xlsx"
    assert convert_excel(str(src), str(out), 100)["tables"] == 3
    col = _text_sheet(out)
    assert "Kurşun kalem" in col and "Silgi" in col
    assert col.index("Ek tablo: bordro özeti") < col.index("Kurşun kalem")


def test_text_row_limit_checked_upfront_in_document_with_tables(invoice: Path, tmp_path: Path, monkeypatch):
    monkeypatch.setattr(excel_mod, "MAX_ROWS", 3)
    out = tmp_path / "o.xlsx"
    with pytest.raises(TooManyRows):
        convert_excel(str(invoice), str(out), 100)
    assert not out.exists()


@pytest.mark.parametrize("target", ["FFFF", "FFFE"])
def test_non_xml_character_cleaned_on_text_sheet_in_document_with_tables(tmp_path: Path, target: str, tr_font: str):
    lines = [f"Merhaba dunya bu bir deneme satiri numara {i} A" for i in range(30)]
    doc = pymupdf.open(_tounicode_pdf(tmp_path / "ffff.pdf", target, lines))
    page = doc.new_page()
    page.insert_font(fontname="tr", fontfile=tr_font)
    tab = [["Ay", "Gelir"], ["Ocak", "12.500,00"]]
    for r, row in enumerate(tab):
        for c, text in enumerate(row):
            rect = pymupdf.Rect(60 + 150 * c, 100 + 24 * r, 210 + 150 * c, 124 + 24 * r)
            page.draw_rect(rect, color=(0, 0, 0), width=0.8)
            page.insert_text((rect.x0 + 4, rect.y1 - 7), text, fontname="tr", fontsize=10)
    src = tmp_path / "mixed-table.pdf"
    doc.save(src)
    doc.close()
    out = tmp_path / "o.xlsx"
    assert convert_excel(str(src), str(out), 100)["tables"] >= 1
    col = _text_sheet(out)
    assert sum(v.endswith("\ufffd") for v in col) == 30
    assert not any(ch in v for v in col for ch in ("\uffff", "\ufffe"))


def test_file_name_never_in_output(invoice: Path, tmp_path: Path):
    out = tmp_path / "o.xlsx"
    convert_excel(str(invoice), str(out), 100)
    with zipfile.ZipFile(out) as z:
        for name in z.namelist():
            assert b"secret-customer-name" not in z.read(name), name


def test_errors_use_contract_classes(tmp_path: Path, invoice: Path):
    broken = tmp_path / "b.pdf"
    broken.write_bytes(b"%PDF-1.7\n broken")
    with pytest.raises(Unreadable):
        convert_excel(str(broken), str(tmp_path / "o.xlsx"), 100)

    with pytest.raises(TooManyPages):
        convert_excel(str(invoice), str(tmp_path / "o.xlsx"), 1)

    encrypted = tmp_path / "s.pdf"
    doc = pymupdf.open(invoice)
    doc.save(encrypted, encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="secret", owner_pw="owner")
    doc.close()
    with pytest.raises(EncryptedPdf):
        convert_excel(str(encrypted), str(tmp_path / "o.xlsx"), 100)

    with pytest.raises(ScannedPdf):
        convert_excel(str(image_only_pdf(tmp_path / "tr.pdf", 3)), str(tmp_path / "o.xlsx"), 100)

    # The classes' exit codes follow roadmap §E (the same contract as test_package from 2b).
    assert [c().exit_code for c in (ScannedPdf, TooManyPages, Unreadable, EncryptedPdf)] == [10, 11, 12, 13]
    assert issubclass(ScannedPdf, ConvertError)


def test_cli_stdout_single_json_line(invoice: Path, tmp_path: Path):
    out = tmp_path / "output.xlsx"
    env = {**os.environ, "PYTHONPATH": SRC}
    p = subprocess.run(
        [sys.executable, "-m", "pdfb_convert", "excel", "--in", str(invoice), "--out", str(out), "--max-pages", "100"],
        capture_output=True,
        text=True,
        check=False,
        env=env,
    )
    assert p.returncode == 0
    lines = p.stdout.strip().splitlines()
    assert len(lines) == 1
    assert json.loads(lines[0]) == {"pages": 2, "tables": 3}
    assert "secret-customer-name" not in p.stdout + p.stderr


def test_cli_excel_extension_required(invoice: Path, tmp_path: Path):
    env = {**os.environ, "PYTHONPATH": SRC}
    p = subprocess.run(
        [
            sys.executable,
            "-m",
            "pdfb_convert",
            "excel",
            "--in",
            str(invoice),
            "--out",
            str(tmp_path / "o.docx"),
            "--max-pages",
            "100",
        ],
        capture_output=True,
        text=True,
        check=False,
        env=env,
    )
    assert p.returncode == 2  # cli.BAD_ARGS


def _run_cli(code: str, *args: str) -> subprocess.CompletedProcess[str]:
    env = {**os.environ, "PYTHONPATH": SRC}
    return subprocess.run(
        [sys.executable, "-c", code, *args], capture_output=True, text=True, check=False, env=env, timeout=240
    )


# Wrapper that patches the converter: the patch is applied inside the converter, after fd isolation (test_cli pattern).
_PATCHED = (
    "import sys\n"
    "from pdfb_convert import cli\n"
    "_, real = cli.CONVERTERS['excel']\n"
    "def patched(src, dst, n):\n"
    "    import pymupdf\n"
    "    original = pymupdf.Page.find_tables\n"
    "    def failing(self, *a, **k):\n"
    "        if self.number == 1:\n"
    "            raise {exc}\n"
    "        return original(self, *a, **k)\n"
    "    pymupdf.Page.find_tables = failing\n"
    "    return real(src, dst, n)\n"
    "cli.CONVERTERS['excel'] = ('.xlsx', patched)\n"
    "raise SystemExit(cli.main(sys.argv[1:]))\n"
)


@pytest.mark.parametrize("name", sorted(BROKEN_TREES))
def test_broken_page_tree_is_unreadable(tmp_path: Path, name: str):
    # A PDF that opens but raises a MuPDF error while pages are read (2b review I2): UNREADABLE_FILE, not
    # CONVERT_FAILED.
    src = tmp_path / f"{name}.pdf"
    src.write_bytes(BROKEN_TREES[name])
    with pytest.raises(Unreadable):
        convert_excel(str(src), str(tmp_path / "o.xlsx"), 100)
    env = {**os.environ, "PYTHONPATH": SRC}
    args = ["excel", "--in", str(src), "--out", str(tmp_path / "o.xlsx"), "--max-pages", "100"]
    p = subprocess.run(
        [sys.executable, "-m", "pdfb_convert", *args], capture_output=True, text=True, check=False, env=env
    )
    assert p.returncode == 12 and p.stdout == ""


def test_late_mupdf_error_is_unreadable(invoice: Path, tmp_path: Path, monkeypatch):
    # A MuPDF error raised while extracting tables (after the scanned check) is also UNREADABLE_FILE (12).
    original = pymupdf.Page.find_tables

    def failing(self, *a, **k):
        if self.number == 1:
            raise pymupdf.FileDataError("corrupt content stream")
        return original(self, *a, **k)

    monkeypatch.setattr(pymupdf.Page, "find_tables", failing)
    out = tmp_path / "o.xlsx"
    with pytest.raises(Unreadable):
        convert_excel(str(invoice), str(out), 100)
    assert not out.exists()
    code = _PATCHED.format(exc="pymupdf.FileDataError('x')")
    p = _run_cli(code, "excel", "--in", str(invoice), "--out", str(out), "--max-pages", "100")
    assert p.returncode == 12 and p.stdout == ""


def test_page_error_is_not_skipped(invoice: Path, tmp_path: Path, monkeypatch):
    # A failing page is not silently skipped to produce an incomplete workbook (2b's ignore_page_error=False rule): a
    # non-MuPDF error is not a ConvertError, and the CLI turns it into 1 (CONVERT_FAILED).
    original = pymupdf.Page.find_tables

    def failing(self, *a, **k):
        if self.number == 1:
            raise KeyError("unexpected")
        return original(self, *a, **k)

    monkeypatch.setattr(pymupdf.Page, "find_tables", failing)
    out = tmp_path / "o.xlsx"
    with pytest.raises(KeyError):
        convert_excel(str(invoice), str(out), 100)
    assert not out.exists()
    for exc in ("KeyError('x')", "MemoryError()", "SystemExit(0)"):
        p = _run_cli(_PATCHED.format(exc=exc), "excel", "--in", str(invoice), "--out", str(out), "--max-pages", "100")
        assert p.returncode == 1 and p.stdout == "", exc


def test_cli_exit_codes(invoice: Path, tmp_path: Path):
    env = {**os.environ, "PYTHONPATH": SRC}
    out = str(tmp_path / "o.xlsx")

    def run(src: Path, max_pages: str = "100") -> subprocess.CompletedProcess[str]:
        args = ["excel", "--in", str(src), "--out", out, "--max-pages", max_pages]
        return subprocess.run(
            [sys.executable, "-m", "pdfb_convert", *args], capture_output=True, text=True, env=env, check=False
        )

    encrypted = tmp_path / "s.pdf"
    with pymupdf.open(invoice) as doc:
        doc.save(encrypted, encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="secret", owner_pw="owner")
    broken = tmp_path / "b.pdf"
    broken.write_bytes(b"%PDF-1.7\n broken")
    cases = [
        (image_only_pdf(tmp_path / "t.pdf", 2), "100", 10),
        (invoice, "1", 11),
        (broken, "100", 12),
        (encrypted, "100", 13),
    ]
    for src, max_pages, code in cases:
        p = run(src, max_pages)
        assert (p.returncode, p.stdout) == (code, ""), src.name


def test_importing_cli_does_not_load_libraries():
    # pymupdf and openpyxl are only loaded after fd isolation (inside the converter).
    code = "import sys; import pdfb_convert.cli; print(sorted(m for m in ('pymupdf', 'openpyxl') if m in sys.modules))"
    p = _run_cli(code)
    assert p.returncode == 0 and p.stdout.strip() == "[]"


def test_workbook_properties_are_neutral(tmp_path: Path, tr_font: str):
    # Counterpart of the DOCX core property cleanup: neither the input's metadata nor openpyxl's "openpyxl" author.
    src = make_excel_invoice(tmp_path / "f.pdf", tr_font)
    tagged = tmp_path / "meta.pdf"
    with pymupdf.open(src) as doc:
        doc.set_metadata({"title": "SECRET-TITLE", "author": "SECRET-AUTHOR", "subject": "SECRET-SUBJECT"})
        doc.save(tagged)
    out = tmp_path / "o.xlsx"
    convert_excel(str(tagged), str(out), 100)
    # load_workbook fills a missing creator back in as "openpyxl", so the file itself is inspected.
    with zipfile.ZipFile(out) as z:
        core = z.read("docProps/core.xml")
        app = z.read("docProps/app.xml")
    for tag in (b"creator", b"lastModifiedBy", b"title", b"subject", b"description", b"keywords", b"category"):
        assert tag not in core, tag
    assert b"openpyxl" not in core
    for leak in (b"SECRET", str(tmp_path).encode(), b"meta.pdf"):
        assert leak not in core + app, leak


def _tounicode_pdf(path: Path, target_hex: str, lines: list[str]) -> Path:
    """PDF whose letter "A" maps to the target_hex code point via ToUnicode (review I3 PoC, ffff.pdf)."""
    cmap = (
        "/CIDInit /ProcSet findresource begin 12 dict begin begincmap /CMapName /X def\n"
        "1 begincodespacerange <00> <FF> endcodespacerange\n"
        f"1 beginbfchar <41> <{target_hex}> endbfchar\n"
        "endcmap CMapName currentdict /CMap defineresource pop end end"
    )
    content = "\n".join(f"BT /F1 12 Tf 72 {750 - 20 * i} Td ({t}) Tj ET" for i, t in enumerate(lines))
    objs = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
        "/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /ToUnicode 6 0 R >>",
        f"<< /Length {len(content)} >>\nstream\n{content}\nendstream",
        f"<< /Length {len(cmap)} >>\nstream\n{cmap}\nendstream",
    ]
    out = b"%PDF-1.7\n"
    offs = []
    for i, o in enumerate(objs, 1):
        offs.append(len(out))
        out += f"{i} 0 obj\n{o}\nendobj\n".encode()
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{o:010d} 00000 n \n".encode() for o in offs)
    out += f"trailer << /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path.write_bytes(out)
    return path


@pytest.mark.parametrize("target", ["FFFF", "FFFE"])
def test_pdf_with_non_xml_characters_converts(tmp_path: Path, target: str):
    # Review I3: a PDF mapping ToUnicode to U+FFFF made the lxml save fail (CONVERT_FAILED); now U+FFFD is written.
    lines = [f"Merhaba dunya bu bir deneme satiri numara {i} A" for i in range(30)]
    src = _tounicode_pdf(tmp_path / "ffff.pdf", target, lines)
    out = tmp_path / "o.xlsx"
    assert convert_excel(str(src), str(out), 100) == {"pages": 1, "tables": 0}
    col = [c.value for c in load_workbook(out)["Metin"]["A"] if c.value is not None]
    assert len(col) == 30
    assert all(v.endswith("\ufffd") for v in col)


def test_app_xml_does_not_reveal_server_software(invoice: Path, tmp_path: Path):
    # Review M8: openpyxl writes "Microsoft Excel Compatible / Openpyxl 3.1.5" and AppVersion into app.xml.
    out = tmp_path / "o.xlsx"
    convert_excel(str(invoice), str(out), 100)
    with zipfile.ZipFile(out) as z:
        app = z.read("docProps/app.xml")
    assert b"openpyxl" not in app.lower()
    assert b"<Application>Microsoft Excel</Application>" in app
    assert b"AppVersion" not in app
    # The patch only lasts for the save: openpyxl's class is restored.
    from openpyxl.packaging import extended  # noqa: PLC0415

    assert extended.ExtendedProperties().Application.startswith("Microsoft Excel Compatible")


def test_text_sheet_row_limit_checked_upfront(tmp_path: Path, tr_font: str, monkeypatch):
    # Review M7: a Metin sheet that would exceed Excel's row limit (1,048,576) is not written; explicit error
    # (CONVERT_FAILED, 1).
    src = make_excel_no_tables(tmp_path / "t.pdf", tr_font)
    monkeypatch.setattr(excel_mod, "MAX_ROWS", 3)
    out = tmp_path / "o.xlsx"
    with pytest.raises(TooManyRows) as info:
        convert_excel(str(src), str(out), 100)
    assert info.value.exit_code == 1 and str(info.value) == ""
    assert not out.exists()
    code = (
        "import sys\n"
        "from pdfb_convert import cli\n"
        "_, real = cli.CONVERTERS['excel']\n"
        "def patched(src, dst, n):\n"
        "    from pdfb_convert import excel\n"
        "    excel.MAX_ROWS = 3\n"
        "    return real(src, dst, n)\n"
        "cli.CONVERTERS['excel'] = ('.xlsx', patched)\n"
        "raise SystemExit(cli.main(sys.argv[1:]))\n"
    )
    p = _run_cli(code, "excel", "--in", str(src), "--out", str(out), "--max-pages", "100")
    assert p.returncode == 1 and p.stdout == ""


def test_truncated_cell_count_goes_to_internal_log(invoice: Path, tmp_path: Path, monkeypatch, caplog):
    # Review M6: the truncated cell count is not part of the stdout contract; it is logged internally as a number only.
    monkeypatch.setattr(xlsx_safe, "MAX_CELL_CHARS", 6)
    out = tmp_path / "o.xlsx"
    with caplog.at_level("INFO", logger="pdfb_convert.excel"):
        assert convert_excel(str(invoice), str(out), 100) == {"pages": 2, "tables": 3}
    records = [r for r in caplog.records if r.name == "pdfb_convert.excel"]
    assert len(records) == 1
    assert records[0].args and isinstance(records[0].args[0], int) and records[0].args[0] > 0
    assert "Danış" not in records[0].getMessage()
    ws = load_workbook(out)["Sayfa 1 - Tablo 1"]
    assert all(len(str(c.value)) <= 6 for row in ws.iter_rows() for c in row if isinstance(c.value, str))
