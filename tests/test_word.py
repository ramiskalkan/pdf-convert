# SPDX-License-Identifier: AGPL-3.0-only

import zipfile
from datetime import UTC, datetime

import pymupdf
import pytest
from docx import Document
from pdf2docx.page.Page import Page

from conftest import BROKEN_TREES, TR_TEXT
from pdfb_convert.errors import ConvertError, ScannedPdf, Unreadable
from pdfb_convert.word import convert_word


def test_converts_turkish_text(make_text_pdf, tmp_path):
    src = make_text_pdf([TR_TEXT, "İkinci sayfa: çğıöşü"])
    dst = tmp_path / "output.docx"
    result = convert_word(str(src), str(dst), 100)
    assert result["pages"] == 2
    text = "\n".join(p.text for p in Document(dst).paragraphs)
    for word in ["Dilekçe", "Işık", "ğüşöç", "İĞÜŞÖÇ", "₺1.234,56", "İkinci"]:
        assert word in text


def test_scanned_pdf_rejected_before_conversion(make_image_pdf, tmp_path):
    dst = tmp_path / "output.docx"
    with pytest.raises(ScannedPdf):
        convert_word(str(make_image_pdf(2)), str(dst), 100)
    assert not dst.exists()


def test_input_path_never_in_output(make_text_pdf, tmp_path):
    # The input's path/name (/job/input.pdf in the sandbox; a distinctive name here) never enters any part of the DOCX.
    src = make_text_pdf(["gizli ad testi metni burada"], name="PERSONAL-NAME-ID.pdf")
    dst = tmp_path / "output.docx"
    convert_word(str(src), str(dst), 100)
    with zipfile.ZipFile(dst) as z:
        for info in z.infolist():
            data = z.read(info)
            assert b"PERSONAL-NAME" not in data, info.filename
            assert str(tmp_path).encode() not in data, info.filename


@pytest.mark.parametrize("name", sorted(BROKEN_TREES))
def test_broken_page_tree_is_unreadable(make_broken_tree, tmp_path, name):
    # A page tree cycle passes opening and FzErrorFormat is raised while pages are read: UNREADABLE_FILE, not
    # CONVERT_FAILED.
    with pytest.raises(Unreadable):
        convert_word(str(make_broken_tree(name)), str(tmp_path / "o.docx"), 100)


def test_page_error_is_not_skipped(make_text_pdf, tmp_path, monkeypatch):
    # By default pdf2docx silently skips a failing page (ignore_page_error=True); an incomplete document is not a
    # success.
    original = Page.parse

    def failing(self, **settings):
        if self.id == 1:
            raise ValueError("broken page")
        return original(self, **settings)

    monkeypatch.setattr(Page, "parse", failing)
    src = make_text_pdf(["birinci sayfa metni yeterince uzun"] * 3)
    with pytest.raises(Exception) as info:  # any conversion error; the CLI turns it into 1
        convert_word(str(src), str(tmp_path / "o.docx"), 100)
    assert not isinstance(info.value, ConvertError)


def test_core_properties_are_neutral(make_text_pdf, tmp_path):
    # DOCX core properties carry neither the input's metadata nor the python-docx template's author/comment/2013 date.
    src = make_text_pdf(["meta veri testi için yeterince uzun metin"])
    tagged = tmp_path / "meta.pdf"
    with pymupdf.open(src) as doc:
        doc.set_metadata({"title": "SECRET-TITLE", "author": "SECRET-AUTHOR", "subject": "SECRET-SUBJECT"})
        doc.save(tagged)
    before = datetime.now(UTC).replace(microsecond=0)
    dst = tmp_path / "output.docx"
    convert_word(str(tagged), str(dst), 100)
    props = Document(dst).core_properties
    for field in ("author", "comments", "title", "subject", "keywords", "category", "last_modified_by"):
        assert getattr(props, field) == "", field
    for stamp in (props.created, props.modified):
        assert stamp is None or (stamp if stamp.tzinfo else stamp.replace(tzinfo=UTC)) >= before
    with zipfile.ZipFile(dst) as z:
        core = z.read("docProps/core.xml")
    for leak in (b"python-docx", b"2013-", b"SECRET"):
        assert leak not in core, leak
