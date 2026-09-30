# SPDX-License-Identifier: AGPL-3.0-only

import io
import random
import zipfile
from datetime import date

import pytest
from openpyxl import Workbook, load_workbook

from pdfb_convert.tr_numbers import Parsed
from pdfb_convert.xlsx_safe import FORMULA_PREFIXES, put_cell, sheet_title

DANGEROUS = [
    "=1+2",
    '=HYPERLINK("http://x")',
    "+90 212 555 00 00",
    "-2+3",
    "@SUM(A1)",
    "\t=1",
    "\r=1",
    "=cmd|' /C calc'!A0",
]


def _save(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_formula_like_text_is_a_text_cell():
    wb = Workbook()
    ws = wb.active
    for i, v in enumerate(DANGEROUS, start=1):
        put_cell(ws, i, 1, v, None)
    data = _save(wb)
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        sheet = z.read("xl/worksheets/sheet1.xml").decode()
    assert "<f>" not in sheet and "<f " not in sheet
    ws2 = load_workbook(io.BytesIO(data)).active
    for i, v in enumerate(DANGEROUS, start=1):
        c = ws2.cell(row=i, column=1)
        assert c.data_type == "s"
        assert c.value == v


def test_negative_number_is_number_formula_is_text():
    wb = Workbook()
    ws = wb.active
    put_cell(ws, 1, 1, "-1.234,56", Parsed(-1234.56, "#,##0.00"))
    put_cell(ws, 2, 1, "-2+3", None)
    put_cell(ws, 3, 1, "25.09.2026", Parsed(date(2026, 9, 25), "dd.mm.yyyy"))
    ws2 = load_workbook(io.BytesIO(_save(wb))).active
    assert ws2["A1"].data_type == "n" and ws2["A1"].value == pytest.approx(-1234.56)
    assert ws2["A1"].number_format == "#,##0.00"
    assert ws2["A2"].data_type == "s" and ws2["A2"].value == "-2+3"
    assert ws2["A3"].is_date and ws2["A3"].number_format == "dd.mm.yyyy"


def test_no_cell_is_a_formula_random():
    rnd = random.Random(7)  # noqa: S311 (deterministic test data, not crypto)
    alphabet = "=+-@\t\r0123456789.,%₺TLabcğüşıöç ()!|'\""
    wb = Workbook()
    ws = wb.active
    for r in range(1, 301):
        v = "".join(rnd.choice(alphabet) for _ in range(rnd.randint(1, 12)))
        put_cell(ws, r, 1, v, None)
    with zipfile.ZipFile(io.BytesIO(_save(wb))) as z:
        assert b"<f" not in z.read("xl/worksheets/sheet1.xml")


def test_control_characters_are_removed():
    wb = Workbook()
    ws = wb.active
    put_cell(ws, 1, 1, "a\x00b\x07c", None)
    assert load_workbook(io.BytesIO(_save(wb))).active["A1"].value == "abc"


def test_empty_cell_is_not_written():
    wb = Workbook()
    ws = wb.active
    put_cell(ws, 1, 1, None, None)
    put_cell(ws, 1, 2, "", None)
    assert ws["A1"].value is None and ws["B1"].value is None


def test_prefix_list_matches_spec():
    assert FORMULA_PREFIXES == ("=", "+", "-", "@", "\t", "\r")


def test_title_rules_and_31_characters():
    used: set[str] = set()
    assert sheet_title("Sayfa 3 - Tablo 1", used) == "Sayfa 3 - Tablo 1"
    long = sheet_title("Sayfa 100 - Tablo 12 [özet]: gelir/gider?*", used)
    assert len(long) <= 31
    for ch in "[]:*?/\\":
        assert ch not in long


def test_title_collision_is_case_insensitive():
    used: set[str] = set()
    a = sheet_title("Metin", used)
    b = sheet_title("METİN", used)
    c = sheet_title("metin", used)
    assert a == "Metin"
    assert len({a.casefold(), b.casefold(), c.casefold()}) == 3


def test_long_title_stays_unique():
    used: set[str] = set()
    base = "Sayfa 100 - Tablo 12 uzun başlık metni"
    names = [sheet_title(base, used) for _ in range(12)]
    assert len(set(n.casefold() for n in names)) == 12
    assert all(len(n) <= 31 for n in names)


def test_openpyxl_accepts_titles():
    used: set[str] = set()
    wb = Workbook()
    for base in ["Sayfa 1 - Tablo 1", "History", "'tırnak'", "a" * 40, "a" * 40]:
        wb.create_sheet(sheet_title(base, used))
    _save(wb)


@pytest.mark.parametrize("bad", ["￾", "￿", "\ud800", "\udfff"])
def test_non_xml_characters_are_replaced(bad: str):
    # Review I3: U+FFFE/U+FFFF (and lone surrogates) do not exist in XML 1.0; lxml would fail the save. They become
    # U+FFFD.
    wb = Workbook()
    ws = wb.active
    put_cell(ws, 1, 1, f"a{bad}b", None)
    assert load_workbook(io.BytesIO(_save(wb))).active["A1"].value == "a�b"


def test_long_text_is_truncated_at_32767_characters():
    # Review M6: an Excel cell holds at most 32,767 characters; more triggers Excel's "repair" warning.
    wb = Workbook()
    ws = wb.active
    assert put_cell(ws, 1, 1, "ş" * 40_000, None) is True
    assert put_cell(ws, 2, 1, "ş" * 32_767, None) is False
    assert put_cell(ws, 3, 1, "kısa", None) is False
    ws2 = load_workbook(io.BytesIO(_save(wb))).active
    assert len(ws2["A1"].value) == 32_767 and len(ws2["A2"].value) == 32_767


@pytest.mark.parametrize("base", ["History", "history", "HISTORY", "  History  "])
def test_reserved_history_title_is_not_used(base: str):
    # Review M10: Excel reserves the title "History" (case-insensitive).
    used: set[str] = set()
    name = sheet_title(base, used)
    assert name.casefold() != "history"
    assert len(name) <= 31


def test_title_never_ends_with_apostrophe():
    # Review M10: Excel rejects ' at the start or end of a title; this is checked after truncation too.
    used: set[str] = set()
    for base in ["x" * 30 + "'yyy", "'a'", "b''", "x" * 26 + "'''" + "zz"]:
        name = sheet_title(base, used)
        assert not name.endswith("'") and not name.startswith("'"), name
        assert 0 < len(name) <= 31
    many = [sheet_title("x" * 27 + "'", used) for _ in range(3)]
    assert all(not n.endswith("'") for n in many)
