# SPDX-License-Identifier: AGPL-3.0-only

"""Corpus criteria (spec §11.3, PDF → Word): expected strings in reading order, table count and cell values."""

import pytest
from docx import Document

from corpus.make_corpus import build
from pdfb_convert.errors import ScannedPdf
from pdfb_convert.word import convert_word


@pytest.fixture(scope="module")
def corpus(tmp_path_factory, tr_font):
    return build(tmp_path_factory.mktemp("corpus"), tr_font)


def _text(path) -> str:
    d = Document(path)
    body = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            body.extend(c.text for c in row.cells)
    return "\n".join(body)


def test_petition_strings_in_reading_order(corpus, tmp_path):
    out = tmp_path / "o.docx"
    convert_word(str(corpus["petition"]), str(out), 100)
    text = _text(out)
    order = ["İSTANBUL", "DAVACI", "Işıklı", "ığdır", "₺12.345,67", "Saygılarımla"]
    positions = [text.find(s) for s in order]
    assert all(p >= 0 for p in positions), dict(zip(order, positions, strict=True))
    assert positions == sorted(positions)


def test_invoice_table_cells(corpus, tmp_path):
    out = tmp_path / "o.docx"
    result = convert_word(str(corpus["invoice-table"]), str(out), 100)
    assert result["tables"] >= 1
    table = Document(out).tables[0]
    header = [c.text.strip() for c in table.rows[0].cells]
    assert header[:5] == ["Sıra", "Ürün", "Miktar", "Birim Fiyat", "Tutar"]
    assert len(table.rows) == 9


def test_multicolumn_keeps_text(corpus, tmp_path):
    out = tmp_path / "o.docx"
    convert_word(str(corpus["newsletter-columns"]), str(out), 100)
    assert "Bülten sütunu" in _text(out)  # layout may differ (decision 4); no text is lost


@pytest.mark.parametrize("name", ["scanned", "mixed-scanned"])
def test_scanned_corpus_rejected(corpus, tmp_path, name):
    # mixed-scanned: 2 of 20 pages have text (90% low-text): even if the browser's sample hits a text page, the server
    # rejects it (Review Focus 1).
    with pytest.raises(ScannedPdf):
        convert_word(str(corpus[name]), str(tmp_path / "o.docx"), 100)


@pytest.mark.slow
def test_hundred_pages_within_limit(corpus, tmp_path):
    result = convert_word(str(corpus["table-100p"]), str(tmp_path / "o.docx"), 100)
    assert result == {"pages": 100, "tables": result["tables"]} and result["tables"] >= 100
